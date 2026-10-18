"""Módulo de Procesamiento Dinámico de Planillas de Laboratorio Vial.

Permite analizar planillas Excel (.xlsx, .xls) y CSV (.csv) extrayendo datos
numéricos y calculando volumetría de mezclas asfálticas (Marshall, Rice, Granulometría)
con precisión matemática determinística antes de enviar el diagnóstico a Gemini.
"""

from __future__ import annotations

import io
from typing import Any, Dict, List, Optional
import pandas as pd


class LabSpreadsheetService:
    """Servicio dinámico de lectura y pre-cálculo de planillas de laboratorio vial."""

    @staticmethod
    def inspect_and_parse(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        """Lee el archivo binario y retorna un resumen estructurado para Gemini."""
        summary: Dict[str, Any] = {
            "filename": filename,
            "sheets": [],
            "preview_tables": [],
            "computed_metrics": {},
        }

        try:
            if filename.lower().endswith((".xlsx", ".xls")):
                excel_file = pd.ExcelFile(io.BytesIO(file_bytes))
                summary["sheets"] = excel_file.sheet_names
                for sheet_name in excel_file.sheet_names[:3]:
                    df = excel_file.parse(sheet_name)
                    df_clean = df.dropna(how="all").dropna(axis=1, how="all")
                    summary["preview_tables"].append({
                        "sheet": sheet_name,
                        "columns": [str(c) for c in df_clean.columns.tolist()[:10]],
                        "row_count": len(df_clean),
                        "head_markdown": df_clean.head(8).to_markdown(index=False),
                    })
            else:
                # Archivo CSV
                df = pd.read_csv(io.BytesIO(file_bytes))
                df_clean = df.dropna(how="all")
                summary["sheets"] = ["CSV_Data"]
                summary["preview_tables"].append({
                    "sheet": "Datos_Principales",
                    "columns": [str(c) for c in df_clean.columns.tolist()[:10]],
                    "row_count": len(df_clean),
                    "head_markdown": df_clean.head(8).to_markdown(index=False),
                })
        except Exception as exc:
            summary["error"] = f"No se pudo parsear como tabla estructurada: {exc}"

        return summary

    @staticmethod
    def calculate_volumetrics(
        gmm: float,
        gmb: float,
        gsb: Optional[float] = None,
        asphalt_pct: Optional[float] = None,
    ) -> Dict[str, float]:
        """Ejecuta cálculos volumétricos determinísticos de mezclas asfálticas."""
        results: Dict[str, float] = {}

        # 1. Vacíos de Aire (%Va)
        if gmm > 0 and gmb > 0:
            va = ((gmm - gmb) / gmm) * 100.0
            results["va_percent"] = round(va, 2)

        # 2. VAM (% Vacíos en el Agregado Mineral)
        if gsb and gsb > 0 and gmb > 0 and asphalt_pct is not None:
            aggregate_pct = 100.0 - asphalt_pct
            vam = 100.0 - ((gmb * aggregate_pct) / gsb)
            results["vam_percent"] = round(vam, 2)

            # 3. VFA (% Vacíos Llenos de Asfalto)
            if "va_percent" in results and vam > results["va_percent"]:
                vfa = ((vam - results["va_percent"]) / vam) * 100.0
                results["vfa_percent"] = round(vfa, 2)

        return results

    @classmethod
    def extract_marshall_parameters(cls, file_bytes: bytes, filename: str) -> Optional[Dict[str, Any]]:
        """Extrae automáticamente variables de ensayo Marshall desde un archivo Excel y calcula volumetría."""
        if not filename.lower().endswith((".xlsx", ".xls")):
            return None

        try:
            import openpyxl
            from services.excel_injector import ExcelInjectorService

            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet = wb["marshall"] if "marshall" in wb.sheetnames else wb.active
            row_map = ExcelInjectorService._find_dynamic_row_mapping(sheet)

            r_seco = row_map.get("peso_seco", 15)
            r_sum = row_map.get("peso_sumergido", 16)
            r_sss = row_map.get("peso_sss", 17)
            r_temp = row_map.get("temperatura", 18)
            r_gmm = row_map.get("Gmm", 19)

            # Buscar valor de Rice / Gmm global en la fila correspondiente
            general_gmm: Optional[float] = None
            for c in range(2, 10):
                val = sheet.cell(r_gmm, c).value
                if isinstance(val, (int, float)) and float(val) > 0:
                    general_gmm = float(val)
                    break

            series: List[Dict[str, Any]] = []
            base_columns = [4, 6, 8, 2, 3, 5]

            for c in base_columns:
                val_seco = sheet.cell(r_seco, c).value
                val_sum = sheet.cell(r_sum, c).value
                val_sss = sheet.cell(r_sss, c).value
                val_temp = sheet.cell(r_temp, c).value
                val_gmm = sheet.cell(r_gmm, c).value

                if (
                    isinstance(val_seco, (int, float))
                    and isinstance(val_sum, (int, float))
                    and isinstance(val_sss, (int, float))
                ):
                    den = float(val_sss) - float(val_sum)
                    gmb = round(float(val_seco) / den, 3) if den > 0 else None
                    gmm_f = float(val_gmm) if isinstance(val_gmm, (int, float)) else general_gmm
                    va = (
                        round(((gmm_f - gmb) / gmm_f) * 100, 2)
                        if (gmb is not None and gmm_f is not None and gmm_f > 0)
                        else None
                    )

                    series.append(
                        {
                            "probeta": len(series) + 1,
                            "peso_seco": float(val_seco),
                            "peso_sumergido": float(val_sum),
                            "peso_sss": float(val_sss),
                            "temperatura": float(val_temp) if isinstance(val_temp, (int, float)) else None,
                            "Gmm": gmm_f,
                            "Gmb": gmb,
                            "Va": va,
                        }
                    )

            if not series:
                return None

            valid_gmbs = [s["Gmb"] for s in series if s["Gmb"] is not None]
            valid_vas = [s["Va"] for s in series if s["Va"] is not None]
            promedio_gmb = round(sum(valid_gmbs) / len(valid_gmbs), 3) if valid_gmbs else None
            promedio_va = round(sum(valid_vas) / len(valid_vas), 2) if valid_vas else None
            first_s = series[0]

            return {
                "series": series,
                "peso_seco": first_s.get("peso_seco"),
                "peso_sumergido": first_s.get("peso_sumergido"),
                "peso_sss": first_s.get("peso_sss"),
                "temperatura": first_s.get("temperatura"),
                "Gmm": first_s.get("Gmm") or general_gmm,
                "Gmb": first_s.get("Gmb"),
                "Va_porcentaje": first_s.get("Va"),
                "promedio_gmb": promedio_gmb,
                "promedio_va": promedio_va,
            }
        except Exception:
            return None
