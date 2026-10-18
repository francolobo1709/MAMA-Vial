"""Servicio de inyección de datos de campo en planillas Excel sin pérdida de formato.

Implementa:
1. Mapeo dinámico por etiquetas de encabezado con fallback determinístico.
2. Inyección multiserie (1, 2 o más probetas con promedios volumétricos).
3. Preservación absoluta de estilos, fuentes, bordes y fórmulas de openpyxl.
"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import openpyxl
from openpyxl.utils import get_column_letter


class ExcelInjectorService:
    """Gestiona la apertura, inyección de valores y guardado de planillas de laboratorio."""

    DEFAULT_TEMPLATE = "12-10-26.xlsx"

    # Etiquetas semánticas para localización dinámica de celdas
    LABEL_PATTERNS = {
        "peso_seco": [r"peso\s*seco", r"seco\s*al\s*aire"],
        "peso_sumergido": [r"peso\s*sumergido", r"sumergido\s*en\s*agua"],
        "peso_sss": [r"peso\s*sss", r"superficie\s*seca", r"saturado.*superficie"],
        "temperatura": [r"temperatura", r"temp.*compactaci"],
        "Gmm": [r"rice", r"gmm", r"te[oó]rica\s*m[aá]xima"],
    }

    # Fallback predeterminado a las celdas estándar de 12-10-26.xlsx
    DEFAULT_COORDS = {
        "peso_seco": (15, 4),        # D15
        "peso_sumergido": (16, 4),   # D16
        "peso_sss": (17, 4),         # D17
        "temperatura": (18, 4),      # D18
        "Gmm": (19, 4),              # D19
    }

    @classmethod
    def _find_dynamic_row_mapping(cls, sheet: Any) -> Dict[str, int]:
        """Escanea las columnas de texto de la planilla para ubicar las filas de cada variable."""
        row_map: Dict[str, int] = {}
        max_scan_row = min(sheet.max_row, 35)

        for r in range(1, max_scan_row + 1):
            for c in range(1, 5):  # Columnas A, B, C, D
                val = str(sheet.cell(row=r, column=c).value or "").strip().lower()
                if not val:
                    continue
                for var_key, patterns in cls.LABEL_PATTERNS.items():
                    if var_key not in row_map:
                        for pat in patterns:
                            if re.search(pat, val, re.IGNORECASE):
                                row_map[var_key] = r
                                break

        # Completar variables no encontradas con el fallback estándar
        for var_key, (default_r, _) in cls.DEFAULT_COORDS.items():
            if var_key not in row_map:
                row_map[var_key] = default_r

        return row_map

    @classmethod
    def inject_field_data(
        cls,
        field_data: Dict[str, Any],
        template_path: str = DEFAULT_TEMPLATE,
        output_path: str = "12-10-26_Actualizada.xlsx",
        sheet_name: str = "marshall",
    ) -> Dict[str, Any]:
        """Inyecta datos de una o múltiples probetas en la planilla conservando fórmulas y estilos."""
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"No se encontró la planilla base '{template_path}'.")

        wb = openpyxl.load_workbook(template_path)
        sheet = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active

        # 1. Localización dinámica de filas
        row_map = cls._find_dynamic_row_mapping(sheet)

        # 2. Normalizar series de probetas (soporte multiserie)
        series_input: List[Dict[str, Any]] = []
        if "series" in field_data and isinstance(field_data["series"], list) and len(field_data["series"]) > 0:
            series_input = field_data["series"]
        else:
            series_input = [field_data]

        injected_records: Dict[str, Dict[str, Any]] = {}
        calculated_series: List[Dict[str, Any]] = []

        # Determinar columnas para inyección (Probeta 1 en col D=4, Probeta 2 en col F=6 si multiserie)
        base_columns = [4, 6, 8]  # D, F, H para no colisionar con unidades

        for idx, specimen in enumerate(series_input):
            col_idx = base_columns[idx] if idx < len(base_columns) else 4 + idx
            col_letter = get_column_letter(col_idx)

            p_seco = specimen.get("peso_seco")
            p_sum = specimen.get("peso_sumergido")
            p_sss = specimen.get("peso_sss")
            temp = specimen.get("temperatura")
            gmm = specimen.get("Gmm") or field_data.get("Gmm")

            specimen_injected: Dict[str, Any] = {}

            var_values = {
                "peso_seco": p_seco,
                "peso_sumergido": p_sum,
                "peso_sss": p_sss,
                "temperatura": temp,
                "Gmm": gmm,
            }

            for var_key, raw_val in var_values.items():
                if raw_val is not None:
                    row_num = row_map.get(var_key, cls.DEFAULT_COORDS[var_key][0])
                    coord = f"{col_letter}{row_num}"
                    try:
                        num_val = float(raw_val)
                        sheet[coord] = num_val
                        specimen_injected[var_key] = {"cell": coord, "value": num_val}
                    except (ValueError, TypeError):
                        sheet[coord] = str(raw_val)
                        specimen_injected[var_key] = {"cell": coord, "value": str(raw_val)}

            injected_records[f"probeta_{idx + 1}"] = specimen_injected

            # Cálculos individuales de volumetría
            gmb_val: Optional[float] = None
            va_val: Optional[float] = None

            try:
                if p_seco is not None and p_sum is not None and p_sss is not None:
                    den = float(p_sss) - float(p_sum)
                    if den > 0:
                        gmb_val = round(float(p_seco) / den, 3)

                if gmb_val is not None and gmm is not None:
                    gmm_f = float(gmm)
                    if gmm_f > 0:
                        va_val = round(((gmm_f - gmb_val) / gmm_f) * 100, 2)
            except Exception:
                pass

            calculated_series.append(
                {
                    "probeta": idx + 1,
                    "peso_seco": p_seco,
                    "peso_sumergido": p_sum,
                    "peso_sss": p_sss,
                    "temperatura": temp,
                    "Gmm": gmm,
                    "Gmb": gmb_val,
                    "Va": va_val,
                }
            )

        wb.save(output_path)

        # Cálculo de promedios para la serie
        valid_gmbs = [s["Gmb"] for s in calculated_series if s["Gmb"] is not None]
        valid_vas = [s["Va"] for s in calculated_series if s["Va"] is not None]

        promedio_gmb = round(sum(valid_gmbs) / len(valid_gmbs), 3) if valid_gmbs else None
        promedio_va = round(sum(valid_vas) / len(valid_vas), 2) if valid_vas else None

        # Para compatibilidad con probeta única
        first_calc = calculated_series[0] if calculated_series else {}

        return {
            "output_path": output_path,
            "sheet_name": sheet.title,
            "injected": injected_records,
            "series": calculated_series,
            "peso_seco": first_calc.get("peso_seco"),
            "peso_sumergido": first_calc.get("peso_sumergido"),
            "peso_sss": first_calc.get("peso_sss"),
            "temperatura": first_calc.get("temperatura"),
            "Gmm": first_calc.get("Gmm"),
            "Gmb": first_calc.get("Gmb"),
            "Va_porcentaje": first_calc.get("Va"),
            "promedio_gmb": promedio_gmb,
            "promedio_va": promedio_va,
        }
